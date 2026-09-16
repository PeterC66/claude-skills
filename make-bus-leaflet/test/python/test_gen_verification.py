"""gen_verification.py -- the document a person reads to decide whether a map is safe.

THE SIXTEENTH OF THE TWENTY, AND THE FIRST WHOSE OUTPUT IS AN ARTEFACT RATHER
THAN A DATA FILE. `verify_report.js` runs the S6 pass and writes
`verification.json`; this module turns that into `verification.docx`, which is
the file Peter opens, the file `stage.js commit S6` puts into `_latest/`, and one
of the two things in an S6 run folder that git tracks. Nothing downstream reads
it -- which is exactly the exposure. A wrong number in a data file eventually
contradicts something; a wrong sentence in a Word document is simply believed,
and the only reader who could notice is the person the document is addressed to.

WHAT THE ESTATE'S OWN REPORTS SAID WHILE THIS FILE DID NOT EXIST. Both faults
below were measured on the committed tree on 2026-09-15, not reasoned about.

  * **The verdict has three values and the module read two.** `summary.verdict`
    is `pass`, `blocked` or `not-verified-uncurated-s1`, and `summary.pass` is
    FALSE for the third -- so an uncurated run printed a red banner reading
    "BLOCKED: hard findings must be resolved" over a run with **no hard findings
    at all**, two lines under a subtitle saying `0 hard`. `verify_report.js`'s
    own console says `NOT VERIFIED` for the same run. No S6 on the estate has
    ever been uncurated, so this one is LATENT and is stated as such.

  * **The qualification the JSON records so it would outlive the console died
    one step later.** `verify_report.js` writes `borrowedRedteam` under a comment
    saying in terms that "the console banner that qualified it does not outlive
    the JSON". **15 S6 runs across nine place maps carry one**; every one is
    verdict `pass`; and every one of their reports printed "RESULT: PASS -- no
    blocking findings; the stored data is safe to build/rely on" with nothing
    beside it. On five of them that pass exists only because a HARD finding was
    restated as a soft. The standing prose was worse than silent -- it asserted
    that "an independent blind red-team agent re-derived **the town's** services
    from scratch", which for a borrowed answer is a claim about a different map.
    This one is not latent. It shipped, fifteen times.

  * And a third, cosmetic and just as invisible: a town's routes.json carries
    `"version": "4.0"` and a place's carries `"version": "v1.0"`, and the
    subtitle printed `routes v` in front of whichever it got. **24 of the 79
    tracked verification.docx read `routes vv1.0`.**

WHY NO EXISTING CHECK COULD SEE ANY OF THEM. The byte gates compare SVG; the
quality ratchet reads metrics; `latest-mirror-gate.js` hashes this docx against
the S6 run's copy of the same docx, so it is exact about WHICH file is mirrored
and says nothing about what is inside it; and `check-s6-claims.mjs` reads the
JSON, never the document. A .docx is the one artefact in the estate whose
correctness has only ever been asserted by somebody opening it.

WHAT IS ASSERTED IS THE RENDERED DOCUMENT, NEVER THE CODE PATH. Every case here
writes a real .docx to a temp directory through `main()` with `sys.argv` patched,
re-opens it with python-docx, and reads the paragraphs and table cells back. That
is deliberate: three of the four helpers in this module manipulate XML directly
(`w:shd`, `w:tblGrid`), so a test that asserted on the objects the module builds
would be asserting on the same mental model that produced the bug.

TWO THINGS ARE DELIBERATELY NOT ASSERTED, and saying so is the point. The column
WIDTHS `set_col_widths` writes are a layout preference with no correct value --
`gen_disagreements.py` carries the same helper and the same reasoning -- so what
is asserted is that the grid and the cells agree with each other, which is the
property the helper's own comment says was the bug. And the created/modified
stamps are asserted only to be later than python-docx's own 2013 default, not to
be any particular time: `datetime.now()` is an input nobody declares, and an
assertion on it would be a clock-dependent artefact of the kind this estate has
already reddened `main` with once.
"""
import contextlib
import io
import json
import os
import shutil
import sys
import tempfile
import unittest

from docx import Document

import _engine

GV = _engine.load("gen_verification")


def a_finding(fid, severity="soft", **kw):
    f = {"id": fid, "severity": severity, "category": kw.get("category", "naming"),
         "message": kw.get("message", "a finding about %s" % fid),
         "source": kw.get("source", "sanity")}
    if "route" in kw:
        f["route"] = kw["route"]
    if "evidence" in kw:
        f["evidence"] = kw["evidence"]
    return f


class Rendered(object):
    """A written .docx, read back as the reader meets it."""

    def __init__(self, path):
        self.path = path
        self.doc = Document(path)
        self.paragraphs = [p.text for p in self.doc.paragraphs if p.text.strip()]
        self.tables = [[[c.text for c in r.cells] for r in t.rows] for t in self.doc.tables]

    @property
    def banner(self):
        return self.tables[0][0][0]

    @property
    def findings_table(self):
        return self.tables[1]

    @property
    def subtitle(self):
        return self.paragraphs[1]

    @property
    def whole(self):
        return "\n".join(self.paragraphs + [c for t in self.tables for r in t for c in r])


class GenVerificationCase(unittest.TestCase):

    def setUp(self):
        self.dir = tempfile.mkdtemp(prefix="gv-")
        self.addCleanup(shutil.rmtree, self.dir, True)

    def render(self, data, out_name=None, argv_out=True):
        """Write `data` as a verification.json and drive main() over it."""
        src = os.path.join(self.dir, "verification.json")
        with open(src, "w", encoding="utf-8") as fh:
            json.dump(data, fh)
        out = os.path.join(self.dir, out_name or "out.docx")
        argv = ["gen_verification.py", src] + ([out] if argv_out else [])
        saved = sys.argv[:]
        sys.argv[:] = argv
        try:
            # main() prints one confirmation line per call; swallowed so a
            # 30-case run does not bury the runner's own output.
            with contextlib.redirect_stdout(io.StringIO()):
                GV.main()
        finally:
            sys.argv[:] = saved
        return Rendered(out if argv_out
                        else os.path.join(self.dir, "verification.docx"))


# =====================================================================
# 1. THE VERDICT -- the one sentence in this document a reader acts on
# =====================================================================
class TheVerdictBanner(GenVerificationCase):

    def test_pass_says_pass(self):
        r = self.render({"town": "T", "generatedAt": "2026-09-15T10:00", "findings": [],
                         "summary": {"hard": 0, "soft": 0, "pass": True, "verdict": "pass"}})
        self.assertIn("RESULT: PASS", r.banner)
        self.assertIn("safe to build", r.banner)

    def test_blocked_says_blocked(self):
        r = self.render({"town": "T", "generatedAt": "2026-09-15T10:00",
                         "findings": [a_finding("F001", "hard")],
                         "summary": {"hard": 1, "soft": 0, "pass": False, "verdict": "blocked"}})
        self.assertIn("RESULT: BLOCKED", r.banner)

    def test_an_uncurated_s1_is_neither_a_pass_nor_a_block(self):
        """THE FAULT. `pass` is false for an uncurated run, so reading it alone
        printed BLOCKED -- "hard findings must be resolved" -- over a run whose
        own subtitle says 0 hard. Latent: no S6 on the estate has been uncurated.
        """
        r = self.render({"town": "T", "generatedAt": "2026-09-15T10:00", "uncuratedS1": True,
                         "findings": [a_finding("F001", "soft")],
                         "summary": {"hard": 0, "soft": 1, "pass": False,
                                     "verdict": "not-verified-uncurated-s1"}})
        self.assertIn("RESULT: NOT VERIFIED", r.banner)
        self.assertNotIn("BLOCKED", r.banner)
        self.assertNotIn("hard findings must be resolved", r.banner)
        self.assertIn("uncurated", r.banner)

    def test_the_uncurated_banner_does_not_depend_on_the_verdict_field(self):
        """19 of the 86 S6 runs on the estate predate `verdict` entirely, so the
        top-level `uncuratedS1` has to be enough on its own."""
        r = self.render({"town": "T", "generatedAt": "2026-09-15T10:00", "uncuratedS1": True,
                         "findings": [], "summary": {"hard": 0, "soft": 0, "pass": False}})
        self.assertIn("RESULT: NOT VERIFIED", r.banner)

    def test_a_file_with_no_summary_at_all_is_judged_by_its_findings(self):
        """The oldest fallback, and the only one that can be wrong quietly: with
        no summary the verdict is derived, so a hard finding must still block."""
        r = self.render({"town": "T", "generatedAt": "2026-09-15T10:00",
                         "findings": [a_finding("F001", "hard")]})
        self.assertIn("RESULT: BLOCKED", r.banner)

    def test_a_file_with_no_summary_and_no_hard_findings_passes(self):
        r = self.render({"town": "T", "generatedAt": "2026-09-15T10:00",
                         "findings": [a_finding("F001", "soft")]})
        self.assertIn("RESULT: PASS", r.banner)

    def test_the_summary_is_authoritative_over_the_findings_it_summarises(self):
        """Not a tautology and worth stating: verify_report.js computes both from
        one array so they cannot part THERE, but the writer knows things the
        array cannot express -- an uncurated S1 is precisely that. A reader
        recomputing the verdict from the rows would lose it."""
        r = self.render({"town": "T", "generatedAt": "2026-09-15T10:00",
                         "findings": [a_finding("F001", "hard")],
                         "summary": {"hard": 1, "soft": 0, "pass": True, "verdict": "pass"}})
        self.assertIn("RESULT: PASS", r.banner)


# =====================================================================
# 2. WHAT QUALIFIES THAT VERDICT -- fifteen shipped runs' worth
# =====================================================================
class TheQualifications(GenVerificationCase):

    BORROWED = {"map": "St Neots", "run": "2026-09-11_1804", "derivedAt": "2026-08-21",
                "borrowedOn": "2026-09-12"}

    def a_borrowed_pass(self, downgraded=None):
        return {"town": "Co-op area", "generatedAt": "2026-09-12T06:17", "redteamPresent": True,
                "findings": [a_finding("F007", "soft", source="redteam")],
                "summary": {"hard": 0, "soft": 1, "pass": True, "verdict": "pass",
                            "borrowedRedteam": self.BORROWED,
                            "downgradedFromHard": downgraded or []}}

    def test_a_borrowed_red_team_is_named_in_the_document(self):
        """THE FAULT, and it shipped. The JSON records this expressly so it
        outlives the console banner; the document a person reads dropped it."""
        r = self.render(self.a_borrowed_pass())
        self.assertIn("This verdict is qualified", r.whole)
        self.assertIn("St Neots", r.whole)
        self.assertIn("borrowed", r.whole.lower())

    def test_a_downgraded_hard_is_counted_where_the_verdict_is(self):
        """Five of the fifteen reached PASS only because a HARD was restated as
        a soft. The explanation was in the finding's own row; the banner said
        "safe to build/rely on" with nothing beside it."""
        r = self.render(self.a_borrowed_pass(downgraded=["F007"]))
        self.assertIn("F007", r.whole)
        self.assertIn("weaker pass", r.whole)

    def test_the_standing_prose_does_not_claim_this_map_was_re_derived(self):
        """It said "re-derived the town's services from scratch". For a borrowed
        answer that is a claim about a different map, in the paragraph that
        explains what the verdict MEANS."""
        r = self.render(self.a_borrowed_pass())
        self.assertNotIn("re-derived the town's services", r.whole)
        self.assertIn("re-derived that map's services", r.whole)

    def test_an_unqualified_pass_carries_no_qualification(self):
        """The control. A sentence printed on every report is a sentence nobody
        reads, and it would make the fifteen above indistinguishable again."""
        r = self.render({"town": "T", "generatedAt": "2026-09-15T10:00", "redteamPresent": True,
                         "findings": [], "summary": {"hard": 0, "soft": 0, "pass": True,
                                                     "verdict": "pass"}})
        self.assertNotIn("This verdict is qualified", r.whole)
        self.assertIn("re-derived the town's services", r.whole)

    def test_an_uncurated_run_says_what_could_not_run(self):
        r = self.render({"town": "T", "generatedAt": "2026-09-15T10:00", "uncuratedS1": True,
                         "findings": [], "summary": {"hard": 0, "soft": 0, "pass": False,
                                                     "verdict": "not-verified-uncurated-s1"}})
        self.assertIn("This verdict is qualified", r.whole)
        self.assertIn("terminus", r.whole)

    def test_no_red_team_at_all_is_said_in_the_prose(self):
        r = self.render({"town": "T", "generatedAt": "2026-09-15T10:00", "redteamPresent": False,
                         "findings": [], "summary": {"hard": 0, "soft": 0, "pass": True,
                                                     "verdict": "pass"}})
        self.assertIn("no red-team file was present", r.whole)

    def test_red_team_sources_are_printed_only_when_there_was_a_red_team(self):
        with_rt = self.render({"town": "T", "generatedAt": "2026-09-15T10:00",
                               "redteamPresent": True, "redteamSources": ["https://a", "https://b"],
                               "findings": [], "summary": {"pass": True, "verdict": "pass"}})
        self.assertIn("https://a; https://b", with_rt.whole)
        without = self.render({"town": "T", "generatedAt": "2026-09-15T10:00",
                               "redteamPresent": False, "redteamSources": ["https://a"],
                               "findings": [], "summary": {"pass": True, "verdict": "pass"}})
        self.assertNotIn("https://a", without.whole)


# =====================================================================
# 3. EVERY FINDING REACHES THE DOCUMENT
# =====================================================================
class EveryFindingIsPrinted(GenVerificationCase):

    def test_hard_findings_come_before_soft_ones_whatever_the_input_order(self):
        r = self.render({"town": "T", "generatedAt": "2026-09-15T10:00",
                         "findings": [a_finding("S1", "soft"), a_finding("H1", "hard"),
                                      a_finding("S2", "soft"), a_finding("H2", "hard")],
                         "summary": {"hard": 2, "soft": 2, "pass": False, "verdict": "blocked"}})
        ids = [row[0] for row in r.findings_table[1:]]
        self.assertEqual(ids, ["H1", "H2", "S1", "S2"])

    def test_a_finding_with_an_unrecognised_severity_is_not_swallowed(self):
        """It was, silently and completely: `hard` and `soft` were two filters
        rather than a partition, so a finding in neither appeared in no row, no
        bullet and no count -- while `if not findings:` stayed False, so even the
        "no findings" row that would have looked odd was suppressed. Latent:
        verify_report.js emits only the two words. A silent drop is worth one
        line to make impossible rather than worth watching for."""
        r = self.render({"town": "T", "generatedAt": "2026-09-15T10:00",
                         "findings": [a_finding("F001", "HARD"), a_finding("F002", "")],
                         "summary": {"hard": 0, "soft": 2, "pass": True, "verdict": "pass"}})
        ids = [row[0] for row in r.findings_table[1:]]
        self.assertIn("F001", ids)
        self.assertIn("F002", ids)

    def test_each_row_carries_the_id_category_route_and_message(self):
        r = self.render({"town": "T", "generatedAt": "2026-09-15T10:00",
                         "findings": [a_finding("F009", "hard", category="terminus", route="401",
                                                message="401 has never gone to Spaldwick")],
                         "summary": {"hard": 1, "soft": 0, "pass": False, "verdict": "blocked"}})
        row = r.findings_table[1]
        self.assertEqual(row[0], "F009")
        self.assertEqual(row[1], "HARD")
        self.assertEqual(row[2], "terminus")
        self.assertEqual(row[3], "401")
        self.assertIn("Spaldwick", row[4])

    def test_the_severity_cell_shouts_only_for_hard(self):
        r = self.render({"town": "T", "generatedAt": "2026-09-15T10:00",
                         "findings": [a_finding("H", "hard"), a_finding("S", "soft")],
                         "summary": {"hard": 1, "soft": 1, "pass": False, "verdict": "blocked"}})
        self.assertEqual([row[1] for row in r.findings_table[1:]], ["HARD", "soft"])

    def test_hard_findings_are_repeated_as_bullets_under_their_own_heading(self):
        r = self.render({"town": "T", "generatedAt": "2026-09-15T10:00",
                         "findings": [a_finding("F001", "hard", message="the blocking one")],
                         "summary": {"hard": 1, "soft": 0, "pass": False, "verdict": "blocked"}})
        self.assertIn("Hard findings", r.whole)
        self.assertTrue(any("the blocking one" in p for p in r.paragraphs))

    def test_a_clean_run_says_so_in_the_table_and_in_the_prose(self):
        r = self.render({"town": "T", "generatedAt": "2026-09-15T10:00", "findings": [],
                         "summary": {"hard": 0, "soft": 0, "pass": True, "verdict": "pass"}})
        self.assertEqual(len(r.findings_table), 2)
        self.assertIn("No findings", r.findings_table[1][4])
        self.assertIn("No findings", r.whole)

    def test_the_evidence_cell_carries_the_source_tag_and_the_evidence(self):
        r = self.render({"town": "T", "generatedAt": "2026-09-15T10:00",
                         "findings": [a_finding("F001", "soft", source="redteam",
                                                evidence={"route": "5", "stops": ["a", "b"]})],
                         "summary": {"hard": 0, "soft": 1, "pass": True, "verdict": "pass"}})
        cell = r.findings_table[1][5]
        self.assertTrue(cell.startswith("[redteam]"))
        self.assertIn("route: 5", cell)
        self.assertIn('stops: ["a","b"]', cell)

    def test_evidence_str_is_empty_for_nothing_and_never_the_word_none(self):
        self.assertEqual(GV.evidence_str(None), "")
        self.assertEqual(GV.evidence_str({}), "")


# =====================================================================
# 4. THE SUBTITLE -- three of its four clauses are conditional
# =====================================================================
class TheSubtitle(GenVerificationCase):

    def test_a_place_version_prints_one_v_and_not_two(self):
        """24 of the 79 tracked verification.docx read `routes vv1.0`, every one
        of them a place map, because a place's routes.json version carries its
        own v and a town's does not."""
        r = self.render({"town": "T", "generatedAt": "2026-09-15T10:00", "findings": [],
                         "inputs": {"routesVersion": "v1.0"},
                         "summary": {"pass": True, "verdict": "pass"}})
        self.assertIn("routes v1.0", r.subtitle)
        self.assertNotIn("vv", r.subtitle)

    def test_a_town_version_still_gains_its_v(self):
        r = self.render({"town": "T", "generatedAt": "2026-09-15T10:00", "findings": [],
                         "inputs": {"routesVersion": "4.0"},
                         "summary": {"pass": True, "verdict": "pass"}})
        self.assertIn("routes v4.0", r.subtitle)

    def test_the_bare_version_helper_touches_nothing_else(self):
        self.assertEqual(GV._bare_version("4.0"), "4.0")
        self.assertEqual(GV._bare_version("v1.0"), "1.0")
        self.assertEqual(GV._bare_version("vv1.0"), "v1.0")   # one v, never a loop
        self.assertEqual(GV._bare_version(3), "3")

    def test_absent_inputs_leave_their_clauses_out_rather_than_printing_none(self):
        r = self.render({"town": "T", "generatedAt": "2026-09-15T10:00", "findings": [],
                         "summary": {"pass": True, "verdict": "pass"}})
        self.assertNotIn("routes", r.subtitle)
        self.assertNotIn("verified", r.subtitle)
        self.assertNotIn("None", r.subtitle)

    def test_the_counts_come_from_the_summary_when_it_has_them(self):
        r = self.render({"town": "T", "generatedAt": "2026-09-15T10:00",
                         "findings": [a_finding("F001", "hard")],
                         "summary": {"hard": 1, "soft": 4, "pass": False, "verdict": "blocked"}})
        self.assertIn("1 hard, 4 soft finding(s)", r.subtitle)

    def test_the_town_and_the_generated_stamp_are_printed_as_given(self):
        r = self.render({"town": "Godmanchester Co-op Ermine Street",
                         "generatedAt": "2026-09-15T10:00", "findings": [],
                         "summary": {"pass": True, "verdict": "pass"}})
        self.assertIn("Godmanchester Co-op Ermine Street", r.paragraphs[0])
        self.assertIn("Generated 2026-09-15T10:00", r.subtitle)


# =====================================================================
# 5. THE FILE IT WRITES, AND THE COMMAND LINE THAT ASKS FOR IT
# =====================================================================
class TheFileAndTheCommandLine(GenVerificationCase):

    def test_the_default_output_lands_beside_the_json_it_was_given(self):
        # Run from a THIRD directory, neither the run folder nor the repository.
        # The mutation this case exists for replaces the joined path with a bare
        # filename, which resolves against the cwd -- so standing in the run
        # folder would make the mutant pass, and standing in the checkout would
        # make it litter the test folder. Both were observed before this line.
        elsewhere = tempfile.mkdtemp(prefix="gv-cwd-")
        self.addCleanup(shutil.rmtree, elsewhere, True)
        here = os.getcwd()
        os.chdir(elsewhere)
        self.addCleanup(os.chdir, here)
        r = self.render({"town": "T", "generatedAt": "2026-09-15T10:00", "findings": [],
                         "summary": {"pass": True, "verdict": "pass"}}, argv_out=False)
        self.assertEqual(os.path.basename(r.path), "verification.docx")
        self.assertEqual(os.path.dirname(os.path.abspath(r.path)), os.path.abspath(self.dir))

    def test_an_explicit_output_path_is_honoured(self):
        r = self.render({"town": "T", "generatedAt": "2026-09-15T10:00", "findings": [],
                         "summary": {"pass": True, "verdict": "pass"}}, out_name="named.docx")
        self.assertTrue(os.path.exists(r.path))

    def test_no_arguments_exits_with_the_usage_rather_than_a_traceback(self):
        saved = sys.argv[:]
        sys.argv[:] = ["gen_verification.py"]
        try:
            with self.assertRaises(SystemExit) as caught:
                GV.main()
        finally:
            sys.argv[:] = saved
        self.assertIn("usage:", str(caught.exception))

    def test_the_created_stamp_is_not_python_docxs_own_2013_default(self):
        """The reason the module stamps them at all: Explorer shows the template's
        2013-12-23 and a reader reads it as the date of the check."""
        r = self.render({"town": "T", "generatedAt": "2026-09-15T10:00", "findings": [],
                         "summary": {"pass": True, "verdict": "pass"}})
        created = r.doc.core_properties.created
        self.assertIsNotNone(created)
        # By YEAR, not by comparison with a constructed datetime: python-docx
        # hands these back timezone-AWARE, and a naive literal beside one is a
        # TypeError rather than a failure -- which reads as a broken test rather
        # than as the stale stamp it is there to catch.
        self.assertGreaterEqual(created.year, 2020)


# =====================================================================
# 6. THE TABLE GEOMETRY -- the helper whose own comment says what the bug was
# =====================================================================
class TheColumnWidths(GenVerificationCase):

    def test_the_grid_and_the_cells_agree(self):
        """`table.autofit` alone is not enough because the PDF is rendered by
        headless LibreOffice, and python-docx leaves `tblGrid` at the equal
        widths it created the table with. The helper rewrites both; that they
        AGREE is the property, not any particular width."""
        from docx.oxml.ns import qn
        r = self.render({"town": "T", "generatedAt": "2026-09-15T10:00",
                         "findings": [a_finding("F001", "soft")],
                         "summary": {"hard": 0, "soft": 1, "pass": True, "verdict": "pass"}})
        table = r.doc.tables[1]
        grid = [int(gc.get(qn("w:w"))) for gc in
                table._tbl.find(qn("w:tblGrid")).findall(qn("w:gridCol"))]
        self.assertEqual(len(grid), 6)
        cells = [c.width.twips for c in table.rows[0].cells]
        self.assertEqual(grid, cells)
        self.assertGreater(grid[4], grid[0])   # Finding is wider than ID

    def test_every_row_has_the_six_columns_the_header_declares(self):
        r = self.render({"town": "T", "generatedAt": "2026-09-15T10:00",
                         "findings": [a_finding("F%03d" % i, "soft") for i in range(5)],
                         "summary": {"hard": 0, "soft": 5, "pass": True, "verdict": "pass"}})
        self.assertTrue(all(len(row) == 6 for row in r.findings_table))


if __name__ == "__main__":
    unittest.main()
